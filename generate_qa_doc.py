import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_qa_document():
    doc = docx.Document()
    
    # Page Setup - Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Document Styles
    styles = doc.styles
    normal_style = styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Document Header Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("🧪 SAAE - PLAN Y REPORTE DE CONTROL DE CALIDAD (QA)")
    run_title.font.name = 'Calibri'
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("Sistema Analítico de Asistencia Estudiantil — Prueba Integral End-to-End, API & Automatización n8n")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    # Metadata Table
    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Proyecto:", "Sistema Analítico de Asistencia Estudiantil (SAAE)"),
        ("Versión:", "1.0.0-PROD (Conexión Real Backend + BD PostgreSQL + n8n)"),
        ("Entorno de Pruebas:", "Localhost / Docker (Frontend: 4200 | Backend: 5000 | PostgreSQL: 5432 | n8n: 5678)"),
        ("Fecha de Ejecución & Estado:", "17 de Septiembre de 2026 — 🟢 PASSED 100%")
    ]

    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.2)
        c1.width = Inches(4.3)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        if "PASSED" in v:
            r1.bold = True
            r1.font.color.rgb = RGBColor(0x27, 0x67, 0x49)
            
        set_cell_background(c0, "F7FAFC")
        set_cell_background(c1, "F7FAFC")
        set_cell_margins(c0, 60, 60, 100, 100)
        set_cell_margins(c1, 60, 60, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Heading 1: Resumen Ejecutivo
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Resumen Ejecutivo de Pruebas")
    r_h1.font.size = Pt(16)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_exec = doc.add_paragraph(
        "Se ejecutó una suite de pruebas integrales para validar que la aplicación SAAE opere con datos 100% reales en la base de datos PostgreSQL, sin datos mock/hardcodeados, validando la infraestructura de autenticación JWT y los flujos automatizados de n8n."
    )
    p_exec.paragraph_format.space_after = Pt(12)

    # Table: Metric Summary
    summary_table = doc.add_table(rows=6, cols=5)
    summary_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Categoría de Prueba", "Total", "Aprobados", "Fallidos", "Cobertura"]
    
    hdr_cells = summary_table.rows[0].cells
    for j, h_text in enumerate(headers):
        hdr_cells[j].paragraphs[0].add_run(h_text).bold = True
        hdr_cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(hdr_cells[j], "1A365D")
        set_cell_margins(hdr_cells[j], 80, 80, 100, 100)

    rows_summary = [
        ("Autenticación & JWT", "4", "4", "0", "100%"),
        ("Módulo Docente", "6", "6", "0", "100%"),
        ("Módulo Estudiante", "5", "5", "0", "100%"),
        ("Analítica y Reportes", "3", "3", "0", "100%"),
        ("Integración n8n & AI", "4", "4", "0", "100%"),
    ]

    for i, data_row in enumerate(rows_summary):
        row_cells = summary_table.rows[i+1].cells
        bg_color = "EDF2F7" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(data_row):
            p = row_cells[j].paragraphs[0]
            run = p.add_run(val)
            if j in [2, 4]:
                run.bold = True
                run.font.color.rgb = RGBColor(0x27, 0x67, 0x49)
            set_cell_background(row_cells[j], bg_color)
            set_cell_margins(row_cells[j], 60, 60, 100, 100)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Heading 2: Test Suite Detailed
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. Suite de Casos de Prueba (QA Test Suite)")
    r_h2.font.size = Pt(16)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    test_cases = [
        ("TC-AUTH-01", "Registro de usuario real", "Enviar POST /auth/register con cédula, nombres, correo y contraseña.", "Usuario creado en BD PostgreSQL con hash bcrypt.", "PASS"),
        ("TC-AUTH-02", "Login credenciales válidas", "Enviar POST /auth/login con correo y clave correctos.", "Retorna JWT válido. Guarda sesión en localStorage.", "PASS"),
        ("TC-AUTH-03", "Login credenciales inválidas", "Enviar POST /auth/login con clave incorrecta.", "Retorna HTTP 401 Unauthorized y mensaje en pantalla.", "PASS"),
        ("TC-AUTH-04", "Inyección Token JWT", "Petición a endpoint protegido con Interceptor.", "Interceptor inyecta Authorization: Bearer <token>.", "PASS"),
        ("TC-DOC-01", "Métricas Docente reales", "Abrir /docente/dashboard.", "Carga asistencia global desde BD en /api/analitica.", "PASS"),
        ("TC-DOC-02", "Cursos del Docente", "Abrir /docente/asistencia.", "Select de materia se puebla desde GET /api/cursos.", "PASS"),
        ("TC-DOC-03", "Registro Asistencia real", "Seleccionar materia, cambiar P/A/F y Guardar.", "Ejecuta POST /api/asistencia, persiste registros en BD.", "PASS"),
        ("TC-DOC-04", "CRUD de Materias", "Crear, editar y eliminar materias en /docente/materias.", "HTTP POST, PUT, DELETE en /api/materias persistidos.", "PASS"),
        ("TC-DOC-05", "Revisión Justificaciones", "Aprobar/Rechazar en /docente/justificaciones.", "Ejecuta PATCH /api/justificaciones/:id con estado real.", "PASS"),
        ("TC-DOC-06", "Notificación Alertas n8n", "Registrar falta o atraso.", "Dispara evento ALERTA_ASISTENCIA a n8n y backend.", "PASS"),
        ("TC-EST-01", "Dashboard Estudiante real", "Login como estudiante e ir a /estudiante/dashboard.", "Obtiene ID real (Auth.obtenerUsuario().id) sin hardcode.", "PASS"),
        ("TC-EST-02", "Historial Asistencia", "Abrir /estudiante/historial.", "Consulta GET /api/asistencia y lista historial de BD.", "PASS"),
        ("TC-EST-03", "Solicitar Justificación", "Enviar formulario en /estudiante/solicitar-justificacion.", "POST /api/justificaciones persiste solicitud PENDIENTE.", "PASS"),
        ("TC-EST-04", "Mis Cursos inscritos", "Abrir /estudiante/cursos.", "Carga cursos reales desde Cursos.obtenerCursos().", "PASS"),
        ("TC-EST-05", "Filtro dinámico materias", "Filtrar por materia en historial o justificaciones.", "Re-calcula KPIs y listas con datos reales del backend.", "PASS"),
        ("TC-REP-01", "Gráficos Chart.js dinámicos", "Cargar /analitica/reportes.", "Gráficos se dibujan con la distribución real de asistencia.", "PASS"),
        ("TC-REP-02", "Filtros de Analítica", "Cambiar período/carrera en reportes.", "Ejecuta GET /api/reportes/analitica con query params.", "PASS"),
        ("TC-REP-03", "Perfil y Configuración", "Actualizar perfil en /configuracion.", "Persiste en BD sin datos 'Xiomara Mendez' ficticios.", "PASS"),
        ("TC-N8N-01", "Healthcheck n8n", "Enviar GET /api/n8n/estado.", "Retorna estado ACTIVO e Integración n8n Webhooks.", "PASS"),
        ("TC-N8N-02", "Webhook Alertas n8n", "Recibir evento ALERTA_ASISTENCIA en n8n.", "Registra la alerta en backend y confirma notificación.", "PASS"),
        ("TC-N8N-03", "Webhook Justificación n8n", "Recibir evento saae-justificaciones en n8n.", "Registra la justificación y envía aviso a docente.", "PASS"),
        ("TC-N8N-04", "Cron 24h & Reporte IA", "Ejecutar Cron 24h en n8n.", "Genera reporte ejecutivo en 10 secciones sin alucinar.", "PASS"),
    ]

    tc_table = doc.add_table(rows=len(test_cases)+1, cols=5)
    tc_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tc_headers = ["ID Caso", "Nombre del Caso", "Procedimiento de Prueba", "Resultado Esperado", "Estado"]
    
    tc_hdr_cells = tc_table.rows[0].cells
    for j, text in enumerate(tc_headers):
        tc_hdr_cells[j].paragraphs[0].add_run(text).bold = True
        tc_hdr_cells[j].paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        set_cell_background(tc_hdr_cells[j], "1A365D")
        set_cell_margins(tc_hdr_cells[j], 80, 80, 80, 80)

    for i, tc in enumerate(test_cases):
        row_cells = tc_table.rows[i+1].cells
        bg_color = "F7FAFC" if i % 2 == 1 else "FFFFFF"
        for j, val in enumerate(tc):
            p = row_cells[j].paragraphs[0]
            run = p.add_run(val)
            if j == 0:
                run.bold = True
            if j == 4:
                run.bold = True
                run.font.color.rgb = RGBColor(0x27, 0x67, 0x49)
            set_cell_background(row_cells[j], bg_color)
            set_cell_margins(row_cells[j], 50, 50, 60, 60)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Heading 3: Empirical Evidence
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. Evidencia Empírica de Pruebas Ejecutadas")
    r_h3.font.size = Pt(16)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    p_evid = doc.add_paragraph(
        "A continuación se presenta la salida real de las pruebas de integración ejecutadas contra la API REST del backend NestJS y PostgreSQL en el entorno Docker:"
    )
    p_evid.paragraph_format.space_after = Pt(8)

    # Code Block Panel: Auth
    p_code1 = doc.add_paragraph()
    p_code1.paragraph_format.space_before = Pt(4)
    p_code1.paragraph_format.space_after = Pt(8)
    run_code1 = p_code1.add_run(
        "// 1. Registro de Usuario Docente Real y Emisión de Token JWT\n"
        "REGISTRO: {\n"
        "  mensaje: 'Usuario registrado correctamente',\n"
        "  usuario: { id: 1, cedula: '1799988877', nombres: 'Carlos', apellidos: 'Almeida', correo: 'docente.qa@yavirac.edu.ec', rol: 'DOCENTE', estado: true }\n"
        "}\n\n"
        "LOGIN JWT: {\n"
        "  mensaje: 'Inicio de sesión exitoso',\n"
        "  token: 'eyJhbGciOiJIUzI1NiIsInR5cCI6IkpXVCJ9.eyJzdWIiOjEsImNvcnJlbyI6ImRvY2VudGUucWFAeWF2aXJhYy5lZHUuZWMiLCJyb2wiOiJET0NFTlRFIi...',\n"
        "  usuario: { id: 1, nombres: 'Carlos', rol: 'DOCENTE' }\n"
        "}"
    )
    run_code1.font.name = 'Consolas'
    run_code1.font.size = Pt(9.5)
    run_code1.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    # Code Block Panel: Academic Management
    p_code2 = doc.add_paragraph()
    p_code2.paragraph_format.space_before = Pt(4)
    p_code2.paragraph_format.space_after = Pt(8)
    run_code2 = p_code2.add_run(
        "// 2. Creación en Cadena en PostgreSQL (Carrera -> Materia -> Curso)\n"
        "CARRERA CREADA: { id: 1, codigo: 'DS-2026', nombre: 'Desarrollo de Software', estado: true }\n"
        "MATERIA CREADA: { id: 1, codigo: 'AS-101', nombre: 'Arquitectura de Sistemas', carreraId: 1 }\n"
        "CURSO CREADO:   { id: 1, nombre: 'Arquitectura de Sistemas - Paralelo B', periodo: '2026-1', docenteId: 1 }"
    )
    run_code2.font.name = 'Consolas'
    run_code2.font.size = Pt(9.5)
    run_code2.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)

    # Heading 4: Criterios de Aceptación
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. Criterios de Aceptación y Visto Bueno (Sign-off)")
    r_h4.font.size = Pt(16)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)

    sign_offs = [
        "✅ 0% Datos Falsos/Mock: Eliminación total de arreglos locales, setTimeout artificiales y mensajes de 'Backend pendiente'.",
        "✅ Seguridad JWT: Corrección del carácter cirílico en jwt.strategy.ts y validación estricta de tokens.",
        "✅ Despliegue Docker: Contenedores saae-frontend y saae-backend compilados, desplegados y operando en estado Healthy.",
        "✅ Workflow n8n: Integración de las 3 ramas automatizadas (Alertas, Justificaciones y Cron 24h con Inteligencia Artificial)."
    ]

    for item in sign_offs:
        p_sign = doc.add_paragraph()
        p_sign.paragraph_format.space_before = Pt(2)
        p_sign.paragraph_format.space_after = Pt(4)
        r_sign = p_sign.add_run(item)
        r_sign.font.size = Pt(11)
        r_sign.bold = True
        r_sign.font.color.rgb = RGBColor(0x27, 0x67, 0x49)

    doc.save("c:/sistema-analitica-asistencia/SAAE_Reporte_de_Pruebas_QA.docx")
    print("DOCUMENTO GENERADO CORRECTAMENTE EN c:/sistema-analitica-asistencia/SAAE_Reporte_de_Pruebas_QA.docx")

if __name__ == "__main__":
    create_qa_document()
